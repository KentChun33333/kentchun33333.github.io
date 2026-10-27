"""Document loaders: PDF/images, Office, CSV/HTML, e-mail, and the source dispatcher."""

from __future__ import annotations

import csv
import email
import hashlib
import io
import re
import zipfile
from email import policy
from pathlib import Path
from typing import Any

import pymupdf

from ...errors import ToolError
from ..session import CHARS_PER_TOKEN
from ..tools import ToolContext
from .common import (MAX_SHEET_ROWS, TEXT_RICH_CHARS, UNIT_TOKENS, Source, Unit, cell, html_text, kind_of,
                     load_image, render, split_text, table_markdown)


def attachment_dir(ctx: ToolContext, path: Path) -> Path:
    digest = hashlib.sha256(path.read_bytes()).hexdigest()[:8]
    return ctx.state_dir / "attachments" / f"{re.sub(r'[^A-Za-z0-9._-]+', '_', path.stem)[:40]}-{digest}"


def save_attachment(folder: Path, name: str, data: bytes) -> Path:
    folder.mkdir(parents=True, exist_ok=True)
    safe = re.sub(r"[^\w.\- ]+", "_", Path(name or "attachment").name).strip() or "attachment"
    target = folder / safe
    n = 1
    while target.exists() and target.read_bytes() != data:
        n += 1
        target = folder / f"{Path(safe).stem}-{n}{Path(safe).suffix}"
    target.write_bytes(data)
    return target


def office_media(ctx: ToolContext, path: Path, src: Source) -> None:
    """Images embedded in .docx/.xlsx/.pptx (word/media/, xl/media/, ppt/media/) are saved as attachments."""
    with zipfile.ZipFile(path) as z:
        media = [n for n in z.namelist() if re.match(r"^(word|xl|ppt)/media/", n)]
        if media:
            folder = attachment_dir(ctx, path)
            src.attachments += [save_attachment(folder, Path(n).name, z.read(n)) for n in media]
            src.notes.append(f"{len(media)} embedded image(s) saved as attachments")


# ---------------------------------------------------------------------- loaders
def load_pdf(ctx: ToolContext, path: Path, pages: str | None) -> Source:
    s = ctx.settings
    vision = ctx.llm.info().vision
    with pymupdf.open(path) as doc:
        numbers = parse_pages(pages, doc.page_count)
        src = Source(path, "pdf", f"{path.name} (PDF, {doc.page_count} pages)")
        xfa = doc.xref_get_key(doc.pdf_catalog(), "AcroForm/XFA")[0] != "null"
        if xfa:
            if s.doc_render == "never" or not vision:
                raise ToolError(f"{path.name}: XFA content not inspected; visual inspection of static pages "
                                "requires rendering and a vision-capable model. Supply a flattened PDF "
                                "exported from an XFA-capable viewer.")
            src.notes.append("XFA form detected: only rendered static pages are inspected; dynamic XFA data "
                             "is not inspected and may be missing. Supply a flattened PDF for complete coverage.")
        if pages:
            src.notes.append(f"pages {pages} only")
        for n in numbers:
            page = doc[n - 1]
            text = page.get_text("text").strip()
            if _pdf_placeholder(text):
                raise ToolError(f"{path.name}, page {n}: PDF viewer placeholder detected; document content "
                                "not inspected. Rendering this page would reproduce the placeholder. "
                                "Supply a flattened PDF exported from a viewer that displays the actual form.")
            mode = s.doc_render
            want = xfa or mode == "always" or (mode == "auto" and
                                             (len(text) < TEXT_RICH_CHARS or has_visuals(page)))
            image = render(page, s.doc_image_px) if want and vision else None
            if xfa:
                src.units.append(Unit(f"page {n}", "", image))
            else:
                src.units += [Unit(f"page {n}", u.text, image if i == 0 else None)
                              for i, u in enumerate(split_text(f"page {n}", text or "[no text layer]"))]
    return src


def _pdf_placeholder(text: str) -> bool:
    text = " ".join(text.casefold().split())
    return (("if this message is not eventually replaced" in text and
             "your pdf viewer may not be able to display" in text) or
            "to view the full contents of this document, you need a later version of the pdf viewer" in text)


def has_visuals(page: pymupdf.Page) -> bool:
    """Figures worth looking at: a raster image over 5% of the page, or many vector drawings (charts, tables)."""
    area = page.rect.width * page.rect.height
    for info in page.get_image_info():
        x0, y0, x1, y1 = info["bbox"]
        if (x1 - x0) * (y1 - y0) > 0.05 * area:
            return True
    return len(page.get_drawings()) >= 20


def parse_pages(spec: str | None, count: int) -> list[int]:
    """'1-5,8' -> [1, 2, 3, 4, 5, 8] (1-based, clipped to the document)."""
    if not spec:
        return list(range(1, count + 1))
    out: list[int] = []
    for part in str(spec).replace(" ", "").split(","):
        if not part:
            continue
        a, _, b = part.partition("-")
        try:
            lo, hi = int(a), int(b) if b else int(a)
        except ValueError:
            raise ToolError(f"pages: cannot read {part!r}; use e.g. '1-5,8'") from None
        out += [n for n in range(lo, hi + 1) if 1 <= n <= count and n not in out]
    if not out:
        raise ToolError(f"pages {spec!r} selects nothing (the document has {count} pages)")
    return out


def load_image_file(ctx: ToolContext, path: Path) -> Source:
    if not ctx.llm.info().vision:
        raise ToolError(f"{ctx.llm.info().id} has no image input; cannot read {path.name}")
    return Source(path, "image", f"{path.name} (image)",
                  [Unit("image", "", load_image(path.read_bytes(), path.name, ctx.settings.doc_image_px))])


def load_docx(ctx: ToolContext, path: Path) -> Source:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = docx.Document(str(path))
    blocks: list[str] = []
    tables = 0
    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, d)
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name or "").lower() if p.style is not None else ""
            if style.startswith("heading"):
                level = int(re.sub(r"\D", "", style) or 1)
                blocks.append("#" * min(level + 1, 6) + " " + text)
            elif style == "title":
                blocks.append("# " + text)
            elif "list" in style:
                blocks.append("- " + text)
            else:
                blocks.append(text)
        elif tag == "tbl":
            tables += 1
            rows = [[cell(c.text) for c in r.cells] for r in Table(child, d).rows]
            blocks.append(f"[table {tables}]\n" + table_markdown(rows))
    props = d.core_properties
    meta = ", ".join(f"{k}: {v}" for k, v in (("author", props.author), ("title", props.title),
                                                ("modified", props.modified)) if v)
    src = Source(path, "docx", f"{path.name} (Word{', ' + meta if meta else ''})")
    src.units = split_text("body", "\n\n".join(blocks) or "[empty document]")
    office_media(ctx, path, src)
    return src


def sheet_units(name: str, rows: list[list[str]]) -> tuple[list[Unit], str | None]:
    while rows and not any(rows[-1]):
        rows.pop()
    note = None
    if len(rows) > MAX_SHEET_ROWS:
        note = f"sheet {name}: {len(rows)} rows, only the first {MAX_SHEET_ROWS} read"
        rows = rows[:MAX_SHEET_ROWS]
    width = max((max((i + 1 for i, v in enumerate(r) if v), default=0) for r in rows), default=0)
    rows = [r[:width] for r in rows]
    if not rows:
        return [Unit(f"sheet {name}", "[empty sheet]")], note
    header = rows[0]
    units, start = [], 0
    limit = UNIT_TOKENS * CHARS_PER_TOKEN
    while start < len(rows):
        size, end = 0, start
        while end < len(rows) and (size < limit or end == start):
            size += sum(len(v) + 3 for v in rows[end])
            end += 1
        block = rows[start:end] if start == 0 else [header] + rows[start:end]
        text = table_markdown(block, 1 if start == 0 else start)
        if start:                               # the repeated header keeps its real row number (1)
            text = text.replace(f"| {start} | ", "| 1 | ", 1)
        units.append(Unit(f"sheet {name} rows {start + 1}-{end}", text))
        start = end
    return units, note


def load_xlsx(ctx: ToolContext, path: Path) -> Source:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    src = Source(path, "xlsx", f"{path.name} (Excel, sheets: {', '.join(wb.sheetnames)})")
    try:
        for ws in wb.worksheets:
            rows = [[cell(v) for v in r] for r in ws.iter_rows(values_only=True)]
            units, note = sheet_units(ws.title, rows)
            src.units += units
            if note:
                src.notes.append(note)
    finally:
        wb.close()
    office_media(ctx, path, src)
    return src


def load_csv(path: Path) -> Source:
    raw = path.read_text(encoding="utf-8-sig", errors="replace")
    dialect = csv.excel_tab if path.suffix.lower() == ".tsv" else csv.Sniffer().sniff(raw[:4096], ",;\t|") \
        if raw.strip() else csv.excel
    rows = [[cell(v) for v in r] for r in csv.reader(io.StringIO(raw), dialect)]
    src = Source(path, "csv", f"{path.name} (table, {len(rows)} rows)")
    src.units, note = sheet_units(path.stem, rows)
    if note:
        src.notes.append(note)
    return src


def load_pptx(ctx: ToolContext, path: Path) -> Source:
    from pptx import Presentation
    prs = Presentation(str(path))
    src = Source(path, "pptx", f"{path.name} (PowerPoint, {len(prs.slides)} slides)")
    for n, slide in enumerate(prs.slides, 1):
        parts = []
        for shape in slide.shapes:
            if getattr(shape, "has_table", False) and shape.has_table:
                parts.append(table_markdown([[cell(c.text) for c in r.cells] for r in shape.table.rows]))
            elif getattr(shape, "has_text_frame", False) and shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            parts.append("Speaker notes: " + slide.notes_slide.notes_text_frame.text.strip())
        src.units += split_text(f"slide {n}", "\n\n".join(parts) or "[no text]")
    office_media(ctx, path, src)
    return src


def email_header(fields: list[tuple[str, Any]]) -> str:
    return "\n".join(f"{k}: {v}" for k, v in fields if v)


def load_eml(ctx: ToolContext, path: Path) -> Source:
    msg = email.message_from_bytes(path.read_bytes(), policy=policy.default)
    head = email_header([("From", msg["from"]), ("To", msg["to"]), ("Cc", msg["cc"]), ("Date", msg["date"]),
                         ("Subject", msg["subject"])])
    body_part = msg.get_body(preferencelist=("plain", "html"))
    body = ""
    if body_part is not None:
        body = body_part.get_content()
        if body_part.get_content_subtype() == "html":
            body = html_text(body)
    src = Source(path, "email", f"{path.name} (e-mail: {msg['subject'] or 'no subject'})")
    folder = attachment_dir(ctx, path)
    for part in msg.iter_attachments():
        payload = part.get_content()
        name = part.get_filename() or ("message.eml" if part.get_content_type() == "message/rfc822" else "attachment")
        data = payload.as_bytes() if hasattr(payload, "as_bytes") else \
            payload.encode() if isinstance(payload, str) else payload
        src.attachments.append(save_attachment(folder, name, data))
    src.units = split_text("email", f"{head}\n\n{body.strip() or '[no body]'}")
    return src


def load_msg(ctx: ToolContext, path: Path) -> Source:
    import extract_msg
    msg = extract_msg.openMsg(str(path))
    try:
        head = email_header([("From", msg.sender), ("To", msg.to), ("Cc", msg.cc), ("Date", msg.date),
                             ("Subject", msg.subject)])
        body = (msg.body or "").strip()
        if not body and msg.htmlBody:
            html = msg.htmlBody.decode("utf-8", "replace") if isinstance(msg.htmlBody, bytes) else msg.htmlBody
            body = html_text(html)
        src = Source(path, "email", f"{path.name} (Outlook e-mail: {msg.subject or 'no subject'})")
        folder = attachment_dir(ctx, path)
        for att in msg.attachments:
            name = att.longFilename or att.shortFilename or "attachment"
            if isinstance(att.data, bytes):
                src.attachments.append(save_attachment(folder, name, att.data))
            elif hasattr(att.data, "export"):            # an attached e-mail: keep it as a .msg of its own
                folder.mkdir(parents=True, exist_ok=True)
                target = folder / (Path(name).stem + ".msg")
                att.data.export(str(target))
                src.attachments.append(target)
            else:
                src.notes.append(f"attachment {name!r} could not be saved")
        src.units = split_text("email", f"{head}\n\n{body or '[no body]'}")
        return src
    finally:
        msg.close()


def load_text(path: Path, kind: str) -> Source:
    raw = path.read_bytes()
    if b"\0" in raw[:8192]:
        raise ToolError(f"{path.name}: binary file of an unsupported type")
    text = raw.decode("utf-8", "replace")
    if kind == "html":
        text = html_text(text)
    return Source(path, kind, f"{path.name} ({kind})", split_text("text", text))


def load_source(ctx: ToolContext, path: Path, pages: str | None = None) -> Source:
    kind = kind_of(path)
    if pages and kind != "pdf":
        raise ToolError("pages applies to PDFs only")
    if kind == "pdf":
        return load_pdf(ctx, path, pages)
    if kind == "image":
        return load_image_file(ctx, path)
    if kind == "docx":
        return load_docx(ctx, path)
    if kind == "xlsx":
        return load_xlsx(ctx, path)
    if kind == "csv":
        return load_csv(path)
    if kind == "pptx":
        return load_pptx(ctx, path)
    if kind == "email":
        return load_eml(ctx, path) if path.suffix.lower() == ".eml" else load_msg(ctx, path)
    return load_text(path, kind)
